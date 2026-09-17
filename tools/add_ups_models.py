#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""增量补充 UPS 配件资料（标准版 + 标准加长版，中英双语）。

数据源：X:/Citops 对外产品资料/1 中文/3 配件/U UPS 与 2 English/3 External accessories/UPS
口径：以中文《MX-UPS 规格书 V5.1》为唯一权威（英文旧规格书写 3000mAh/74x18.8x70mm，
      与中文标准版 3350mAh/99.7x48x79mm 不符，属旧版内容；英文规格书按 V5.1 翻译重做）。

产物：
  data/accessories/media/{zh,en}/accessory/<id>/spec.docx + images/NN.webp [+ spec-images/]
  catalog.json / catalog.enriched.json 中 4 个条目的封面/画廊/型号/参数/摘要

缓存规避：/assets/* 为一年 immutable，已存在的图片/规格书 URL 不得换内容，
         故 acc-en-ups 的新图用 11/12/13.webp 新文件名，规格书用 spec-v5.1.docx。

之后跑：node tools/convert_specs.js --only=<ids> && node tools/build.js
"""
import json
import os
import shutil
import zipfile

import pypdfium2 as pdfium
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image

SRC_CN = "X:/Citops 对外产品资料/1 中文/3 配件/U UPS"
SRC_EN = "X:/Citops 对外产品资料/2 English/3 External accessories/UPS"
ACC = os.path.join("data", "accessories")
MEDIA = os.path.join(ACC, "media")
TMP = ".tmp_ups_build"

# ---------------------------------------------------------------- 源文件
CN_MASTER = os.path.join(SRC_CN, "MX-UPS 规格书 V5.1.docx")          # 权威主规格书
CN_EXT = os.path.join(SRC_CN, "加长版/MX-UPS-07 规格书（加长版）.docx")
IMG_MAIN = os.path.join(SRC_CN, "UPS 外观图.png")                     # MX-UPS-05 外观
IMG_TOPTAIL = os.path.join(SRC_CN, "俯视+尾线（标准版）.png")           # 俯视+尾线
IMG_EXT_PHOTO = os.path.join(SRC_CN, "加长版/加长版照片.jpg")           # MX-UPS-07 外观
PDF_DIM_STD = os.path.join(SRC_CN, "参考资料/外置UPS-外形尺寸图-PDF-B.pdf")
PDF_DIM_EXT = os.path.join(SRC_CN, "加长版/尺寸图.pdf")

# 英文页沿用语言无关的实拍/图纸（EN 自带照片印的是旧型号 MX-UPS-03，不用）


def compress_image(src, dst, max_side=1280, quality=82):
    im = Image.open(src)
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        im = im.convert("RGBA")
    else:
        im = im.convert("RGB")
    w, h = im.size
    if max(w, h) > max_side:
        s = max_side / max(w, h)
        im = im.resize((int(w * s), int(h * s)), Image.LANCZOS)
    im.save(dst, "WEBP", quality=quality)


def render_pdf_webp(pdf_path, dst, scale=2.2, max_side=1600, quality=85):
    pdf = pdfium.PdfDocument(pdf_path)
    page = pdf[0]
    img = page.render(scale=scale).to_pil().convert("RGB")
    pdf.close()
    w, h = img.size
    if max(w, h) > max_side:
        s = max_side / max(w, h)
        img = img.resize((int(w * s), int(h * s)), Image.LANCZOS)
    # 图纸为白底细线，用白底画布保证透明区域可读
    canvas = Image.new("RGB", img.size, "white")
    canvas.paste(img, (0, 0))
    canvas.save(dst, "WEBP", quality=quality)


def extract_docx_media(docx_path, member, dst):
    with zipfile.ZipFile(docx_path) as z, z.open(member) as f, open(dst, "wb") as out:
        shutil.copyfileobj(f, out)


def render_pdf_png(pdf_path, dst, scale=2.6):
    """渲染 PDF 首页为 PNG（供英文规格书内嵌）。"""
    pdf = pdfium.PdfDocument(pdf_path)
    img = pdf[0].render(scale=scale).to_pil().convert("RGB")
    pdf.close()
    canvas = Image.new("RGB", img.size, "white")
    canvas.paste(img, (0, 0))
    canvas.save(dst)


# ---------------------------------------------------------------- EN 规格书
EN_SPEC_ROWS = [
    ("Case material", "Aluminium alloy", ""),
    ("Power input", "DC 8-36V", ""),
    ("Battery type", "18650 lithium battery pack", ""),
    ("Battery capacity",
     "MX-UPS-05: 3350mAh (49.6Wh) / MX-UPS-07: 6700mAh (99.2Wh)",
     "Optional 2600mAh, 15000mAh, etc."),
    ("Battery rating voltage", "14.8V", "4 x 3.70V"),
    ("End-of-charge voltage", "16.8V", "4 x 4.20V"),
    ("Discharge cut-off voltage", "11V", "4 x 2.75V"),
    ("Charging current", "2A", ""),
    ("Charging time",
     "3350mAh: 1.6h +/-10% (<=2h) / 6700mAh: 3.5h +/-10% (<=4h)",
     "Depends on battery capacity version"),
    ("Max output current", "7A", ""),
    ("Dimensions",
     "Standard 99.7 x 48 x 79mm +/-2mm / Extended 140.2 x 48 x 79mm +/-2mm", ""),
    ("Weight", "Standard approx. 700g +/-10% / Extended approx. 1400g +/-10%", ""),
    ("Waterproof grade", "IP66", ""),
    ("Ambient temperature",
     "Charge 0~45C / Discharge -20~60C / Storage -20~50C",
     "Do not install the equipment near a heat source"),
    ("Indicator lights",
     "PWR (red): on when input power present, off when no input; "
     "CHG (green): fast blink = trickle charge, solid = constant-current charge, "
     "off = fully charged", ""),
    ("Input voltage",
     "10-19V step-up charging (92% efficiency) / 19-36V step-down charging (96% efficiency)",
     "Automatic switching; use 16AWG (1.3mm2) wire for the input"),
    ("Output voltage",
     "Input 0-9V: switch to battery voltage 12-16.5V / Input 9-36V: output follows input 8-36V",
     "Battery takes over when input drops below 9V"),
    ("Runtime example",
     "Standard 49.6Wh / 18W = 2.76h, Extended 99.2Wh / 18W = 5.51h",
     "MDVR load at 12V / 1.5A (18W)"),
]

EN_MODEL_MATRIX = [
    ("MX-UPS-03", "3000mAh", "Previous version"),
    ("MX-UPS-03-01", "2600mAh", "Previous version"),
    ("MX-UPS-04", "6000mAh", "Previous version"),
    ("MX-UPS-05", "3350mAh", "Standard"),
    ("MX-UPS-06", "15000mAh", "Special order"),
    ("MX-UPS-07", "6700mAh", "Standard extended"),
]

EN_TAIL_VERSIONS = [
    ("V1.0", "5-pin aviation plug (input 150mm) / 5-pin aviation socket (output 200mm)"),
    ("V2.0", "6-pin square connector (input 150mm / output 200mm)"),
    ("V3.0", "9-pin aviation plug (input 150mm) / 9-pin aviation socket (output 200mm)"),
    ("V4.0", "DJ7031-6.3-11 (input) / DJ7031-6.3-21 (output); RS485 optional, 180mm, DJ7031-2.8-21"),
    ("V5.0", "DJ7031-6.3-21 (input 150mm / output 200mm)"),
]

EN_WIRE_COLORS = [
    ("Input", "Red", "Power input +"),
    ("Input", "Black", "Power input -"),
    ("Input", "Yellow", "ACC ignition"),
    ("Input", "Green", "Not connected"),
    ("Output", "Red", "Not connected"),
    ("Output", "Black", "Power output -"),
    ("Output", "Yellow", "ACC ignition"),
    ("Output", "Green", "Power output +"),
]


def build_en_spec_docx(dst_path, variant, dim_img):
    """按中文 V5.1 内容生成英文规格书（variant: 'std' | 'ext'）。"""
    vname = "Standard (MX-UPS-05, 3350mAh)" if variant == "std" else "Extended (MX-UPS-07, 6700mAh)"
    doc = Document()
    doc.add_heading("Waterproof UPS Specification - %s" % vname, level=0)

    doc.add_heading("Product introduction", level=1)
    doc.add_paragraph(
        "The UPS is an on-board uninterruptible power supply that uses a lithium battery pack to "
        "generate 8-36V DC power for vehicle equipment. It is mainly used to keep vehicle "
        "monitoring equipment powered in special circumstances."
    )
    doc.add_heading("Product advantages", level=1)
    for t in ("Sealed aluminium case, waterproof and dustproof",
              "Automatically supplies power from its own battery when the vehicle power is cut off",
              "Power supply duration can be configured together with the MDVR host"):
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("Model naming", level=1)
    doc.add_paragraph(
        "Model format: MX-UPS-XX-Vx.x. Example: MX-UPS-05-V1.0 = 3350mAh battery capacity, "
        "tail cable version 1.0 (5-pin aviation connector). MX-UPS-07-V3.0 = 6700mAh battery "
        "capacity, tail cable version 3.0 (9-pin aviation connector)."
    )

    doc.add_heading("Core parameters", level=1)
    tb = doc.add_table(rows=1, cols=3)
    tb.style = "Table Grid"
    for i, h in enumerate(("Item", "Parameter", "Note")):
        tb.rows[0].cells[i].text = h
    for a, b, c in EN_SPEC_ROWS:
        row = tb.add_row().cells
        row[0].text, row[1].text, row[2].text = a, b, c

    doc.add_heading("Battery capacity model matrix", level=1)
    tb2 = doc.add_table(rows=1, cols=3)
    tb2.style = "Table Grid"
    for i, h in enumerate(("Model", "Battery capacity", "Version")):
        tb2.rows[0].cells[i].text = h
    for a, b, c in EN_MODEL_MATRIX:
        row = tb2.add_row().cells
        row[0].text, row[1].text, row[2].text = a, b, c

    doc.add_heading("Tail cable versions", level=1)
    tb3 = doc.add_table(rows=1, cols=2)
    tb3.style = "Table Grid"
    tb3.rows[0].cells[0].text = "Version"
    tb3.rows[0].cells[1].text = "Connector"
    for a, b in EN_TAIL_VERSIONS:
        row = tb3.add_row().cells
        row[0].text, row[1].text = a, b

    doc.add_heading("Wire color definition", level=1)
    tb4 = doc.add_table(rows=1, cols=3)
    tb4.style = "Table Grid"
    for i, h in enumerate(("Cable", "Color", "Definition")):
        tb4.rows[0].cells[i].text = h
    for a, b, c in EN_WIRE_COLORS:
        row = tb4.add_row().cells
        row[0].text, row[1].text, row[2].text = a, b, c

    doc.add_heading("Wiring", level=1)
    doc.add_paragraph(
        "Option 1: connect PWR+ together with ACC and control it through the ignition switch. "
        "When ACC is off the UPS switches to battery power for the peripherals; when ACC is on the "
        "vehicle battery supplies the peripherals and the UPS charges its battery."
    )
    doc.add_paragraph(
        "Option 2: connect PWR+ / PWR- to the vehicle battery terminals and ACC to the ignition "
        "switch. When ACC is off the load switches to the UPS battery until the battery voltage "
        "drops below the vehicle battery voltage. This option is only suitable for 12V vehicles - "
        "configure the ACC delayed shutdown time according to the charging and discharging time so "
        "that the vehicle battery is never drained."
    )
    doc.add_paragraph("9-pin aviation connector pinout:")
    doc.add_picture(os.path.join(TMP, "pinout.png"), width=Inches(2.6))
    doc.add_paragraph("Dimension drawing (%s):" % ("standard" if variant == "std" else "extended"))
    doc.add_picture(dim_img, width=Inches(5.6))

    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.size is None:
                r.font.size = Pt(10.5)
    doc.save(dst_path)


# ---------------------------------------------------------------- 条目定义
# (id, lang, name, modelCode, spec_src|None, image_specs, summary, keySpecs, chips)
# image_specs: ('img', path) | ('pdf', path) | ('new', path)
def entries():
    zh_common_specs = [
        ("电源输入", "DC 8-36V（10-19V 升压充电 92%，19-36V 降压充电 96%，自动切换）"),
        ("电池类型", "18650 锂电池组，额定 14.8V（4×3.70V）"),
        ("充电 / 放电电压", "充电结束 16.8V / 放电截止 11V"),
        ("充电电流", "2A"),
        ("最大输出电流", "7A"),
        ("防水等级", "IP66（密封铝合金外壳）"),
        ("工作温度", "充电 0~45℃ / 放电 -20~60℃ / 存储 -20~50℃"),
        ("尾线版本", "V1.0 5芯航空头 / V2.0 6PIN方口 / V3.0 9芯航空头 / V4.0、V5.0 DJ7031 端子（V4.0 可选 RS485）"),
    ]
    en_common_specs = [
        ("Power input", "DC 8-36V (10-19V step-up 92%, 19-36V step-down 96%, auto switching)"),
        ("Battery type", "18650 lithium pack, rated 14.8V (4 x 3.70V)"),
        ("Charge / cut-off voltage", "End of charge 16.8V / discharge cut-off 11V"),
        ("Charging current", "2A"),
        ("Max output current", "7A"),
        ("Waterproof grade", "IP66 (sealed aluminium case)"),
        ("Ambient temperature", "Charge 0-45C / Discharge -20-60C / Storage -20-50C"),
        ("Tail cable versions", "V1.0 5-pin aviation / V2.0 6-pin square / V3.0 9-pin aviation / V4.0 & V5.0 DJ7031 terminals (V4.0 RS485 optional)"),
    ]
    return [
        dict(
            id="acc-zh-U31", lang="zh", type="accessory", group="U",
            name="外置 UPS 电源（标准版）",
            model_code="MX-UPS-05-V1.0",
            spec_src=CN_MASTER, spec_name="spec.docx",
            images=[("img", IMG_MAIN), ("img", IMG_TOPTAIL), ("pdf", PDF_DIM_STD)],
            summary=(
                "车载外置不间断电源，铝合金密封外壳达 IP66 防水防尘。输入 DC 8-36V 宽压，车辆断电后"
                "自动切换电池供电，配合主机可自行设定供电时长。标准版电池容量 3350mAh（49.6Wh），"
                "最大输出 7A，12V/1.5A 负载约可续航 2.76 小时；另有 2600/3000/6000/6700/15000mAh "
                "容量档位与 5 种尾线接口可选。"
            ),
            key_specs=[
                ("产品型号", "MX-UPS-05（3350mAh 标准版）"),
                ("电池容量", "3350mAh / 49.6Wh（可选 2600mAh、15000mAh 等）"),
            ] + zh_common_specs + [
                ("尺寸 / 重量", "99.7×48×79mm ±2mm / 约 700g ±10%"),
                ("充电时间", "约 1.6h±10%（＜2h）"),
                ("放电时长", "MDVR 12V/1.5A（18W）负载约 2.76h"),
            ],
            chips=["3350mAh / 49.6Wh", "DC 8-36V", "IP66 防水", "7A 输出"],
        ),
        dict(
            id="acc-zh-U32", lang="zh", type="accessory", group="U",
            name="外置 UPS 电源（加长版）",
            model_code="MX-UPS-07-V1.0",
            spec_src=CN_EXT, spec_name="spec.docx",
            images=[("img", IMG_EXT_PHOTO), ("pdf", PDF_DIM_EXT)],
            summary=(
                "标准加长版车载 UPS，与标准版同系列铝合金密封外壳（IP66）、同规格输入输出。"
                "电池容量 6700mAh（99.2Wh），尺寸 140.2×48×79mm、约 1400g，12V/1.5A 负载下"
                "约可续航 5.51 小时，适合需要更长断电续航的车载监控场景。"
            ),
            key_specs=[
                ("产品型号", "MX-UPS-07（6700mAh 标准加长版）"),
                ("电池容量", "6700mAh / 99.2Wh"),
            ] + zh_common_specs + [
                ("尺寸 / 重量", "140.2×48×79mm ±2mm / 约 1400g ±10%"),
                ("充电时间", "约 3.5h±10%（＜4h）"),
                ("放电时长", "MDVR 12V/1.5A（18W）负载约 5.51h"),
            ],
            chips=["6700mAh / 99.2Wh", "DC 8-36V", "IP66 防水", "续航 5.5h"],
        ),
        dict(
            id="acc-en-ups", lang="en", type="accessory", group="",
            name="Waterproof UPS (Standard)",
            model_code="MX-UPS-05-V1.0",
            spec_src=None, spec_name="spec-v5.1.docx", en_variant="std",
            images=[("img", IMG_MAIN), ("img", IMG_TOPTAIL), ("pdf", PDF_DIM_STD)],
            image_start=11,
            summary=(
                "On-board uninterruptible power supply with a sealed aluminium case rated IP66. "
                "Wide 8-36V DC input, automatic switch-over to battery power when the vehicle supply "
                "is cut off, and configurable power supply duration with the MDVR host. Standard "
                "version: 3350mAh (49.6Wh), 7A max output, about 2.76h runtime at 12V/1.5A. Other "
                "capacities (2600/3000/6000/6700/15000mAh) and five tail-cable versions are available."
            ),
            key_specs=[
                ("Model", "MX-UPS-05 (3350mAh standard)"),
                ("Battery capacity", "3350mAh / 49.6Wh (2600mAh, 15000mAh optional)"),
            ] + en_common_specs + [
                ("Dimensions / weight", "99.7 x 48 x 79mm +/-2mm / approx. 700g +/-10%"),
                ("Charging time", "approx. 1.6h +/-10% (<=2h)"),
                ("Runtime example", "about 2.76h with a 12V/1.5A (18W) MDVR load"),
            ],
            chips=["3350mAh / 49.6Wh", "DC 8-36V", "IP66", "7A output"],
        ),
        dict(
            id="acc-en-ups-07", lang="en", type="accessory", group="",
            name="Waterproof UPS (Extended)",
            model_code="MX-UPS-07-V1.0",
            spec_src=None, spec_name="spec.docx", en_variant="ext",
            images=[("img", IMG_EXT_PHOTO), ("pdf", PDF_DIM_EXT)],
            summary=(
                "Standard extended version of the on-board UPS with the same sealed IP66 aluminium "
                "case and identical input/output ratings. 6700mAh (99.2Wh) battery, 140.2 x 48 x 79mm, "
                "approx. 1400g, and about 5.51h runtime at 12V/1.5A - suited to vehicle surveillance "
                "installations that need a longer backup time."
            ),
            key_specs=[
                ("Model", "MX-UPS-07 (6700mAh standard extended)"),
                ("Battery capacity", "6700mAh / 99.2Wh"),
            ] + en_common_specs + [
                ("Dimensions / weight", "140.2 x 48 x 79mm +/-2mm / approx. 1400g +/-10%"),
                ("Charging time", "approx. 3.5h +/-10% (<=4h)"),
                ("Runtime example", "about 5.51h with a 12V/1.5A (18W) MDVR load"),
            ],
            chips=["6700mAh / 99.2Wh", "DC 8-36V", "IP66", "5.5h backup"],
        ),
    ]


def main():
    os.makedirs(TMP, exist_ok=True)
    # 9 芯航空头针脚图（语言无关）供英文规格书使用
    extract_docx_media(CN_MASTER, "word/media/image8.png", os.path.join(TMP, "pinout.png"))
    # 尺寸图渲染为 PNG，供英文规格书内嵌
    render_pdf_png(PDF_DIM_STD, os.path.join(TMP, "dim-std.png"))
    render_pdf_png(PDF_DIM_EXT, os.path.join(TMP, "dim-ext.png"))

    cat_path = os.path.join(ACC, "catalog.json")
    enr_path = os.path.join(ACC, "catalog.enriched.json")
    with open(cat_path, encoding="utf-8") as f:
        cat = json.load(f)
    with open(enr_path, encoding="utf-8") as f:
        enr = json.load(f)
    ids_cat = {it["id"] for it in cat["items"]}
    ids_enr = {it["id"] for it in enr["items"]}
    assert ids_cat == ids_enr, "catalog / enriched id 集合不一致"

    for p in (cat_path, enr_path):
        b = p + ".bak-ups"
        if not os.path.exists(b):
            shutil.copy2(p, b)
            print("backup:", b)

    for spec in entries():
        iid, lang = spec["id"], spec["lang"]
        dest = os.path.join(MEDIA, lang, "accessory", iid)
        os.makedirs(os.path.join(dest, "images"), exist_ok=True)

        # ---- 规格书
        spec_rel = "media/%s/accessory/%s/%s" % (lang, iid, spec["spec_name"])
        spec_abs = os.path.join(dest, spec["spec_name"])
        if spec.get("en_variant"):
            build_en_spec_docx(spec_abs, spec["en_variant"],
                               os.path.join(TMP, "dim-%s.png" % spec["en_variant"]))
        else:
            shutil.copy2(spec["spec_src"], spec_abs)
        print("  spec ->", spec_rel)

        # ---- 图片
        start = spec.get("image_start", 1)
        gallery = []
        for i, (kind, src) in enumerate(spec["images"]):
            n = start + i
            dst = os.path.join(dest, "images", "%02d.webp" % n)
            if kind == "pdf":
                render_pdf_webp(src, dst)
            else:
                compress_image(src, dst)
            rel = os.path.relpath(dst, ACC).replace("\\", "/")
            gallery.append(rel)
        cover = gallery[0]

        base = {
            "id": iid, "type": "accessory", "lang": lang, "name": spec["name"],
            "line": "", "model": "", "group": spec["group"],
            "specDocx": spec_rel, "docCount": 1,
            "cover": cover, "gallery": gallery, "imageCount": len(gallery),
        }
        rich = dict(base)
        rich.update({
            "modelCode": spec["model_code"],
            "summary": spec["summary"],
            "keySpecs": [{"k": k, "v": v} for k, v in spec["key_specs"]],
            "chips": spec["chips"],
            "lineKey": base.get("lineKey", ""),
        })

        for book in (cat, enr):
            items = book["items"]
            if iid in {it["id"] for it in items}:
                # 就地更新：保留原 id/顺序，覆盖业务字段
                tgt = next(it for it in items if it["id"] == iid)
                tgt.update(rich if book is enr else base)
            else:
                items.append(rich if book is enr else base)
        print("OK %-16s images=%d cover=%s" % (iid, len(gallery), cover))

    with open(cat_path, "w", encoding="utf-8") as f:
        json.dump(cat, f, ensure_ascii=False, indent=2)
    with open(enr_path, "w", encoding="utf-8") as f:
        json.dump(enr, f, ensure_ascii=False, indent=1)
    print("WROTE catalog items=%d / enriched items=%d" % (len(cat["items"]), len(enr["items"])))


if __name__ == "__main__":
    main()
